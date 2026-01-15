"""
Output Formatter and Download Handler
Formats rewritten content and creates downloadable Word documents
"""

import os
from pathlib import Path
from typing import Dict, Any, Optional
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class OutputFormatter:
    """Formats and exports content to various formats"""

    def __init__(self, output_dir: str = None):
        """
        Initialize output formatter.

        Args:
            output_dir: Directory for output files (uses Config.OUTPUT_DIR if not provided)
        """
        from config import Config
        self.output_dir = Path(output_dir or Config.OUTPUT_DIR)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def format_and_save(
        self,
        stage2_result: Dict[str, Any],
        filename: str = None,
        format_type: str = "docx"
    ) -> str:
        """
        Format and save the rewritten content.

        Args:
            stage2_result: Complete output from Stage 2 rewriting
            filename: Output filename (auto-generated if not provided)
            format_type: Output format ('docx', 'txt', 'md')

        Returns:
            Path to the saved file
        """
        # Generate filename if not provided
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            topic = stage2_result["metadata"].get("article_topic", "article")
            safe_topic = "".join(c for c in topic if c.isalnum() or c in (' ', '-', '_'))[:50]
            filename = f"{safe_topic}_{timestamp}"

        # Add extension if not present
        if not filename.endswith(f".{format_type}"):
            filename = f"{filename}.{format_type}"

        output_path = self.output_dir / filename

        # Format and save based on type
        if format_type == "docx":
            self._save_as_docx(stage2_result, output_path)
        elif format_type == "txt":
            self._save_as_txt(stage2_result, output_path)
        elif format_type == "md":
            self._save_as_markdown(stage2_result, output_path)
        else:
            raise ValueError(f"Unsupported format: {format_type}")

        print(f"Content saved to: {output_path}")
        return str(output_path)

    def _save_as_docx(self, stage2_result: Dict[str, Any], output_path: Path) -> None:
        """Save content as a formatted Word document"""
        doc = Document()

        # Set document properties
        metadata = stage2_result["metadata"]
        doc.core_properties.title = metadata.get("title", "SEO Article")
        doc.core_properties.author = "AI Document Processor"
        doc.core_properties.comments = f"Primary Keyword: {metadata.get('primary_keyword')}"

        # Add title
        title = doc.add_heading(metadata.get("title", "Article Title"), level=0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Add metadata section
        meta_para = doc.add_paragraph()
        meta_para.add_run("Article Topic: ").bold = True
        meta_para.add_run(f"{metadata.get('article_topic')}\n")
        meta_para.add_run("Primary Keyword: ").bold = True
        meta_para.add_run(f"{metadata.get('primary_keyword')}\n")
        meta_para.add_run("Target Word Count: ").bold = True
        meta_para.add_run(f"{metadata.get('target_word_count')}\n")

        if metadata.get("keywords_used"):
            meta_para.add_run("Keywords: ").bold = True
            meta_para.add_run(f"{', '.join(metadata['keywords_used'])}\n")

        # Add separator
        doc.add_paragraph("_" * 80)

        # Add main content
        content = stage2_result["rewritten_content"]
        self._parse_and_add_content(doc, content)

        # Add footer with generation info
        doc.add_page_break()
        footer_para = doc.add_paragraph("\n")
        footer_para.add_run("Document Generation Information").bold = True
        footer_para = doc.add_paragraph()
        footer_para.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")

        if stage2_result.get("stage1_metadata"):
            s1_meta = stage2_result["stage1_metadata"]
            footer_para.add_run(f"Source Documents: {s1_meta.get('num_documents', 'N/A')}\n")
            footer_para.add_run(f"Total Source Words: {s1_meta.get('total_words', 'N/A')}\n")

        # Save document
        doc.save(str(output_path))

    def _parse_and_add_content(self, doc: Document, content: str) -> None:
        """Parse markdown-like content and add to document with formatting"""
        lines = content.split('\n')
        i = 0

        while i < len(lines):
            line = lines[i].strip()

            if not line:
                i += 1
                continue

            # Handle headings
            if line.startswith('#'):
                level = len(line) - len(line.lstrip('#'))
                heading_text = line.lstrip('#').strip()
                doc.add_heading(heading_text, level=min(level, 9))

            # Handle lists
            elif line.startswith(('- ', '* ', '+ ')):
                list_items = [line[2:].strip()]
                i += 1
                while i < len(lines) and lines[i].strip().startswith(('- ', '* ', '+ ')):
                    list_items.append(lines[i].strip()[2:].strip())
                    i += 1
                i -= 1

                for item in list_items:
                    para = doc.add_paragraph(item, style='List Bullet')

            # Handle numbered lists
            elif line and line[0].isdigit() and '. ' in line[:4]:
                list_items = [line.split('. ', 1)[1] if '. ' in line else line]
                i += 1
                while i < len(lines) and lines[i].strip() and lines[i].strip()[0].isdigit():
                    list_items.append(lines[i].strip().split('. ', 1)[1] if '. ' in lines[i] else lines[i].strip())
                    i += 1
                i -= 1

                for item in list_items:
                    para = doc.add_paragraph(item, style='List Number')

            # Regular paragraph
            else:
                para = doc.add_paragraph(line)

            i += 1

    def _save_as_txt(self, stage2_result: Dict[str, Any], output_path: Path) -> None:
        """Save content as plain text"""
        content = stage2_result["rewritten_content"]
        metadata = stage2_result["metadata"]

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"Title: {metadata.get('title')}\n")
            f.write(f"Topic: {metadata.get('article_topic')}\n")
            f.write(f"Primary Keyword: {metadata.get('primary_keyword')}\n")
            f.write("=" * 80 + "\n\n")
            f.write(content)

    def _save_as_markdown(self, stage2_result: Dict[str, Any], output_path: Path) -> None:
        """Save content as markdown"""
        content = stage2_result["rewritten_content"]
        metadata = stage2_result["metadata"]

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"---\n")
            f.write(f"title: {metadata.get('title')}\n")
            f.write(f"topic: {metadata.get('article_topic')}\n")
            f.write(f"keyword: {metadata.get('primary_keyword')}\n")
            f.write(f"date: {datetime.now().strftime('%Y-%m-%d')}\n")
            f.write(f"---\n\n")
            f.write(content)

    def create_analysis_report(self, stage1_result: Dict[str, Any], filename: str = None) -> str:
        """
        Create a report of Stage 1 analysis results.

        Args:
            stage1_result: Complete output from Stage 1 analysis
            filename: Output filename

        Returns:
            Path to the saved report
        """
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"analysis_report_{timestamp}.docx"

        output_path = self.output_dir / filename

        doc = Document()
        doc.add_heading('Document Analysis Report', 0)

        # Metadata
        if "metadata" in stage1_result:
            doc.add_heading('Source Documents', level=1)
            meta = stage1_result["metadata"]
            para = doc.add_paragraph()
            para.add_run(f"Number of Documents: {meta.get('num_documents')}\n")
            para.add_run(f"Total Words: {meta.get('total_words')}\n")
            para.add_run("\nDocument Names:\n")
            for name in meta.get('document_names', []):
                doc.add_paragraph(name, style='List Bullet')

        # Keywords
        if "keywords" in stage1_result:
            doc.add_heading('Suggested Keywords', level=1)
            for i, keyword in enumerate(stage1_result["keywords"], 1):
                doc.add_paragraph(f"{i}. {keyword}", style='List Number')

        # Themes
        if "themes" in stage1_result:
            doc.add_heading('Identified Themes', level=1)
            for theme in stage1_result["themes"]:
                doc.add_heading(theme.get("name", "Theme"), level=2)
                doc.add_paragraph(theme.get("description", ""))

        # NLP Entities
        if "nlp_entities" in stage1_result:
            doc.add_heading('NLP Entities', level=1)
            entities = stage1_result["nlp_entities"]

            doc.add_heading('Key Nouns', level=2)
            doc.add_paragraph(", ".join(entities.get("nouns", [])))

            doc.add_heading('Action Verbs', level=2)
            doc.add_paragraph(", ".join(entities.get("verbs", [])))

            doc.add_heading('Descriptive Adjectives', level=2)
            doc.add_paragraph(", ".join(entities.get("adjectives", [])))

        doc.save(str(output_path))
        print(f"Analysis report saved to: {output_path}")
        return str(output_path)
