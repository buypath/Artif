"""
Document Parser for extracting text from Word documents (.docx)
"""

import os
from typing import List, Dict, Any
from pathlib import Path
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


class DocumentParser:
    """Parse Word documents and extract structured content"""

    def __init__(self, max_size_mb: int = 10):
        """
        Initialize document parser.

        Args:
            max_size_mb: Maximum allowed file size in megabytes
        """
        self.max_size_mb = max_size_mb
        self.supported_extensions = ['.docx']

    def parse_document(self, file_path: str) -> Dict[str, Any]:
        """
        Parse a single Word document.

        Args:
            file_path: Path to the Word document

        Returns:
            Dictionary containing parsed content and metadata
        """
        file_path = Path(file_path)

        # Validate file
        self._validate_file(file_path)

        # Parse document
        doc = Document(str(file_path))

        # Extract content
        content = {
            "filename": file_path.name,
            "filepath": str(file_path),
            "title": self._extract_title(doc),
            "paragraphs": self._extract_paragraphs(doc),
            "tables": self._extract_tables(doc),
            "full_text": self._extract_full_text(doc),
            "metadata": self._extract_metadata(doc),
            "word_count": self._count_words(doc),
        }

        return content

    def parse_multiple_documents(self, file_paths: List[str]) -> List[Dict[str, Any]]:
        """
        Parse multiple Word documents.

        Args:
            file_paths: List of paths to Word documents

        Returns:
            List of parsed document dictionaries
        """
        parsed_docs = []

        for file_path in file_paths:
            try:
                parsed_doc = self.parse_document(file_path)
                parsed_docs.append(parsed_doc)
            except Exception as e:
                print(f"Error parsing {file_path}: {str(e)}")
                # Continue with other documents

        return parsed_docs

    def consolidate_documents(self, parsed_docs: List[Dict[str, Any]]) -> str:
        """
        Consolidate multiple parsed documents into a single formatted text.

        Args:
            parsed_docs: List of parsed document dictionaries

        Returns:
            Consolidated text with document separators
        """
        consolidated_parts = []

        for i, doc in enumerate(parsed_docs, 1):
            doc_section = f"""
{'='*80}
DOCUMENT {i}: {doc['filename']}
{'='*80}

Title: {doc['title']}
Word Count: {doc['word_count']}

{doc['full_text']}
"""
            consolidated_parts.append(doc_section)

        return "\n\n".join(consolidated_parts)

    def _validate_file(self, file_path: Path) -> None:
        """Validate file exists, has correct extension, and size"""
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if file_path.suffix.lower() not in self.supported_extensions:
            raise ValueError(
                f"Unsupported file format: {file_path.suffix}. "
                f"Supported formats: {', '.join(self.supported_extensions)}"
            )

        file_size_mb = file_path.stat().st_size / (1024 * 1024)
        if file_size_mb > self.max_size_mb:
            raise ValueError(
                f"File size ({file_size_mb:.2f}MB) exceeds maximum "
                f"allowed size ({self.max_size_mb}MB)"
            )

    def _extract_title(self, doc: Document) -> str:
        """Extract document title (first heading or first paragraph)"""
        # Try to get the first heading
        for paragraph in doc.paragraphs:
            if paragraph.style.name.startswith('Heading'):
                return paragraph.text.strip()

        # Fall back to first non-empty paragraph
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                return paragraph.text.strip()

        return "Untitled Document"

    def _extract_paragraphs(self, doc: Document) -> List[Dict[str, str]]:
        """Extract all paragraphs with their styles"""
        paragraphs = []

        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append({
                    "text": para.text.strip(),
                    "style": para.style.name,
                })

        return paragraphs

    def _extract_tables(self, doc: Document) -> List[List[List[str]]]:
        """Extract all tables from the document"""
        tables = []

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            tables.append(table_data)

        return tables

    def _extract_full_text(self, doc: Document) -> str:
        """
        Extract all text from document including tables.
        Preserves document structure.
        """
        text_parts = []

        # Iterate through document body in order
        for element in doc.element.body:
            if isinstance(element, CT_P):
                # Paragraph
                paragraph = Paragraph(element, doc)
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())

            elif isinstance(element, CT_Tbl):
                # Table
                table = Table(element, doc)
                table_text = self._format_table_as_text(table)
                if table_text:
                    text_parts.append(table_text)

        return "\n\n".join(text_parts)

    def _format_table_as_text(self, table: Table) -> str:
        """Format table as readable text"""
        table_lines = []

        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                table_lines.append(row_text)

        if table_lines:
            return "[TABLE]\n" + "\n".join(table_lines) + "\n[/TABLE]"
        return ""

    def _extract_metadata(self, doc: Document) -> Dict[str, Any]:
        """Extract document metadata"""
        core_props = doc.core_properties

        metadata = {
            "author": core_props.author or "Unknown",
            "created": str(core_props.created) if core_props.created else None,
            "modified": str(core_props.modified) if core_props.modified else None,
            "last_modified_by": core_props.last_modified_by or "Unknown",
            "revision": core_props.revision,
        }

        return metadata

    def _count_words(self, doc: Document) -> int:
        """Count total words in document"""
        full_text = self._extract_full_text(doc)
        return len(full_text.split())
